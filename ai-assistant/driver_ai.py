#!/usr/bin/env python3
"""
DRIVER AI assistant: ask questions about your crash data in plain English.

A small local language model (through Ollama) turns the question into a JSON
"query plan". This program checks the plan, turns it into SQL and runs it as a
database user that can only read. The model never sees your records and cannot
change anything.

    driver-ai setup                 create the read-only user and views (again after form changes)
    driver-ai dictionary            show what the model is told about your data
    driver-ai ask "question"        ask one question from the terminal
    driver-ai chat                  a conversation in the terminal (follow-up questions)
    driver-ai serve                 start the web page at http://localhost:8800
    driver-ai test                  run a set of test questions

Settings (environment variables, all optional):
    DRIVER_DB_HOST      database address        (default 192.168.12.101)
    DRIVER_AI_MODEL     Ollama model            (default qwen2.5:3b)
    OLLAMA_URL          Ollama address          (default http://127.0.0.1:11434)
    DRIVER_AI_PORT      web page port           (default 8800)
"""
import csv
import datetime as dt
import decimal
import difflib
import io
import json
import os
import re
import secrets
import sys
import time
import urllib.error
import urllib.request
import uuid

try:
    import psycopg2
    from psycopg2 import sql as pgsql
except ImportError:  # pragma: no cover
    sys.exit("psycopg2 is missing. Install it with:\n"
             "  uv pip install --python ~/driver-ai-venv/bin/python psycopg2-binary openpyxl")

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(HERE)
SECRETS_FILE = os.path.join(HERE, ".ai_db.json")
LOG_DIR = os.path.join(HERE, "logs")

DB_HOST = os.environ.get("DRIVER_DB_HOST", "192.168.12.101")
DB_PORT = int(os.environ.get("DRIVER_DB_PORT", "5432"))
DB_NAME = os.environ.get("DRIVER_DB_NAME", "driver")
DB_ADMIN_USER = os.environ.get("DRIVER_DB_ADMIN_USER", "driver")
DB_ADMIN_PASSWORD = os.environ.get("DRIVER_DB_ADMIN_PASSWORD", "driver")
AI_USER = "driver_ai"
RECORD_TYPE_LABEL = os.environ.get("DRIVER_RECORD_TYPE", "Incident")
MODEL = os.environ.get("DRIVER_AI_MODEL", "qwen2.5:3b")
OLLAMA_URL = os.environ.get("OLLAMA_URL", "http://127.0.0.1:11434").rstrip("/")
PORT = int(os.environ.get("DRIVER_AI_PORT", "8800"))
DRIVER_URL = os.environ.get("DRIVER_URL", "http://localhost:7000")

PREVIEW_ROWS = 200        # rows shown in the popup
EXPORT_ROWS = 100000      # maximum rows in an export
MAX_CATEGORIES = 40       # text columns with up to this many values are listed in the dictionary


def local_timezone():
    """Time zone from DRIVER's settings (deployment/ansible/group_vars/development)."""
    try:
        with open(os.path.join(REPO, "deployment", "ansible", "group_vars", "development")) as f:
            m = re.search(r"^local_time_zone_id:\s*['\"]?([\w/+-]+)", f.read(), re.M)
            if m:
                return m.group(1)
    except OSError:
        pass
    return "Asia/Dhaka"


TZ = os.environ.get("DRIVER_TZ") or local_timezone()


class PlanError(Exception):
    """The question could not be turned into a valid query."""


def slug(text):
    s = re.sub(r"[^a-z0-9]+", "_", (text or "").lower()).strip("_")
    if not s:
        s = "field"
    if s[0].isdigit():
        s = "f_" + s
    return s[:50]


def unique(name, used):
    base, n = name, 2
    while name in used:
        name = "%s_%d" % (base, n)
        n += 1
    used.add(name)
    return name


# =================================================================== setup
FUNCTIONS_SQL = """
CREATE OR REPLACE FUNCTION ai.txt(j jsonb) RETURNS text LANGUAGE plpgsql IMMUTABLE AS $$
BEGIN
    IF j IS NULL OR jsonb_typeof(j) = 'null' THEN RETURN NULL; END IF;
    IF jsonb_typeof(j) = 'array' THEN
        RETURN (SELECT string_agg(e, ', ') FROM jsonb_array_elements_text(j) AS e);
    END IF;
    RETURN NULLIF(j #>> '{}', '');
END $$;

CREATE OR REPLACE FUNCTION ai.num(j jsonb) RETURNS numeric LANGUAGE plpgsql IMMUTABLE AS $$
BEGIN
    RETURN (j #>> '{}')::numeric;
EXCEPTION WHEN others THEN
    RETURN NULL;
END $$;
"""


def admin_connect():
    try:
        return psycopg2.connect(host=DB_HOST, port=DB_PORT, dbname=DB_NAME, user=DB_ADMIN_USER,
                                password=DB_ADMIN_PASSWORD, connect_timeout=10)
    except psycopg2.OperationalError as e:
        sys.exit(db_help(e))


def db_help(e):
    msg = "Cannot connect to DRIVER's database at %s:%s.\n%s\n" % (DB_HOST, DB_PORT, str(e).strip())
    if "refused" in str(e):
        msg += ("\nThe database VM's firewall only lets DRIVER's own VMs in. Allow this computer once with:\n"
                "  cd ~/DRIVER\n"
                "  vagrant ssh database -c \"sudo ufw allow proto tcp from 192.168.12.1 to any port 5432\"\n"
                "If DRIVER is not running, start it first: bash start-driver.sh")
    elif "timeout" in str(e).lower() or "No route" in str(e):
        msg += "\nIs DRIVER running? Start it with: bash start-driver.sh"
    return msg


def lit(value):
    """SQL string literal (for identifiers that come from the database itself)."""
    return "'" + str(value).replace("'", "''") + "'"


def field_expr(src, name, prop):
    """SQL expression for one form field inside a JSON object."""
    j = "%s->%s" % (src, lit(name))
    if prop.get("type") in ("number", "integer"):
        return "ai.num(%s)" % j
    return "ai.txt(%s)" % j


def build_views(cur):
    """Generate the CREATE VIEW statements from DRIVER's current form."""
    cur.execute("""SELECT uuid FROM grout_recordtype WHERE label = %s AND active
                   ORDER BY created DESC LIMIT 1""", (RECORD_TYPE_LABEL,))
    row = cur.fetchone()
    if not row:
        sys.exit("No active record type called '%s' found in DRIVER." % RECORD_TYPE_LABEL)
    rt = str(row[0])
    cur.execute("""SELECT schema FROM grout_recordschema WHERE record_type_id = %s AND next_version_id IS NULL
                   ORDER BY version DESC LIMIT 1""", (rt,))
    schema = cur.fetchone()[0]
    defs = schema.get("definitions", {})
    details_key = next((k for k, d in defs.items() if d.get("details")), None)
    if details_key is None:
        details_key = next((k for k, d in defs.items() if not d.get("multiple")), None)

    cur.execute("""SELECT uuid, label, display_field FROM grout_boundary
                   WHERE status IN ('COMPLETE', 'WARNING') AND display_field IS NOT NULL ORDER BY label""")
    boundaries = cur.fetchall()

    base_from = """FROM grout_record r
    JOIN grout_recordschema s ON s.uuid = r.schema_id
    LEFT JOIN data_driverrecord d ON d.record_ptr_id = r.uuid
    WHERE s.record_type_id = %s AND NOT r.archived""" % lit(rt)
    local = "(r.occurred_from AT TIME ZONE %s)" % lit(TZ)
    statements, comments = [], []

    # ---- crashes: one row per record
    used = set()
    cols = [
        ("record_id", "r.uuid", "Record ID"),
        ("occurred_at", local, "Date and time (local)"),
        ("date", local + "::date", "Date"),
        ("year", "EXTRACT(year FROM %s)::int" % local, "Year"),
        ("month", "EXTRACT(month FROM %s)::int" % local, "Month number 1-12"),
        ("year_month", "to_char(%s, 'YYYY-MM')" % local, "Year and month"),
        ("weekday", "trim(to_char(%s, 'Day'))" % local, "Day of the week"),
        ("hour", "EXTRACT(hour FROM %s)::int" % local, "Hour of the day 0-23"),
    ]
    for uid, label, display in boundaries:
        name = slug(label)
        cols.append((name, """(SELECT p.data->>%s FROM grout_boundarypolygon p
            WHERE p.boundary_id = %s AND ST_Within(r.geom, p.geom) LIMIT 1)""" % (lit(display), lit(uid)),
                     "Boundary: " + label))
    if details_key:
        for fname, prop in defs[details_key].get("properties", {}).items():
            if fname == "_localId" or prop.get("media") or prop.get("watch"):
                continue
            cols.append((slug(fname), field_expr("r.data->" + lit(details_key), fname, prop), fname))
    cols += [
        ("road", "d.road", "Road"),
        ("city", "d.city", "City / place"),
        ("address", "r.location_text", "Address"),
        ("weather", "d.weather", "Weather"),
        ("light", "d.light", "Light conditions"),
        ("lat", "round(ST_Y(r.geom)::numeric, 6)", "Latitude"),
        ("lon", "round(ST_X(r.geom)::numeric, 6)", "Longitude"),
    ]
    sections = []
    vused = {"crashes", "blackspots"}
    for key, d in defs.items():
        if key == details_key or not d.get("multiple"):
            continue
        fields = [(n, p) for n, p in d.get("properties", {}).items() if n != "_localId" and not p.get("media")]
        if not fields:
            continue
        view = unique(slug(d.get("plural_title") or d.get("title") or key), vused)
        sections.append((key, d, view, fields))
        cols.append(("n_" + view, "CASE WHEN jsonb_typeof(r.data->%s) = 'array' THEN jsonb_array_length(r.data->%s) "
                     "ELSE 0 END" % (lit(key), lit(key)), "Number of %s listed" % (d.get("plural_title") or view)))
    select = []
    for name, expr, title in cols:
        name = unique(name, used)
        select.append("    %s AS %s" % (expr, name))
        comments.append(("crashes", name, title))
    statements.append("CREATE VIEW ai.crashes AS SELECT\n%s\n%s" % (",\n".join(select), base_from))
    statements.append("COMMENT ON VIEW ai.crashes IS 'One row per crash'")

    # ---- one view per "multiple" section (vehicles, people, ...)
    view_by_key = {key: view for key, _, view, _ in sections}
    for key, d, view, fields in sections:
        used = {"record_id", "item_id"}
        select = ["    r.uuid AS record_id", "    e->>'_localId' AS item_id"]
        for fname, prop in fields:
            target = (prop.get("watch") or {}).get("target")
            if target:
                tview = view_by_key.get(target) or next(
                    (v for k, v in view_by_key.items() if k.lower() == str(target).lower()), None)
                name = unique(slug(fname) + "_item", used)
                select.append("    e->>%s AS %s" % (lit(fname), name))
                comments.append((view, name, "link:" + (tview or "")))
                continue
            name = unique(slug(fname), used)
            select.append("    %s AS %s" % (field_expr("e", fname, prop), name))
            comments.append((view, name, fname))
        statements.append("""CREATE VIEW ai.%s AS SELECT
%s
FROM grout_record r
    JOIN grout_recordschema s ON s.uuid = r.schema_id,
    LATERAL jsonb_array_elements(CASE WHEN jsonb_typeof(r.data->%s) = 'array'
                                      THEN r.data->%s ELSE '[]'::jsonb END) AS e
WHERE s.record_type_id = %s AND NOT r.archived""" % (view, ",\n".join(select), lit(key), lit(key), lit(rt)))
        statements.append("COMMENT ON VIEW ai.%s IS %s" % (
            view, lit("One row per %s entry of a crash (%s)" % (d.get("title") or view, key))))

    # ---- black spots (latest calculation)
    bcols = ["    b.uuid AS blackspot_id",
             "    round(b.severity_score::numeric, 4) AS severity_score",
             "    b.num_records AS num_crashes",
             "    b.num_severe AS num_severe_crashes",
             """    (SELECT d2.road FROM grout_record r2 JOIN data_driverrecord d2 ON d2.record_ptr_id = r2.uuid
        JOIN grout_recordschema s2 ON s2.uuid = r2.schema_id
        WHERE s2.record_type_id = %s AND NOT r2.archived AND d2.road IS NOT NULL AND ST_Within(r2.geom, b.geom)
        GROUP BY d2.road ORDER BY count(*) DESC LIMIT 1) AS road""" % lit(rt)]
    bcomments = [("severity_score", "Risk score (higher = more dangerous)"),
                 ("num_crashes", "Crashes in the black spot"),
                 ("num_severe_crashes", "Severe crashes in the black spot"), ("road", "Main road")]
    bused = {"blackspot_id", "severity_score", "num_crashes", "num_severe_crashes", "road"}
    for uid, label, display in boundaries:
        name = unique(slug(label), bused)
        bcols.append("""    (SELECT p.data->>%s FROM grout_boundarypolygon p
        WHERE p.boundary_id = %s AND ST_Within(ST_Centroid(b.geom), p.geom) LIMIT 1) AS %s""" % (
            lit(display), lit(uid), name))
        bcomments.append((name, "Boundary: " + label))
    bcols += ["    round(ST_Y(ST_Centroid(b.geom))::numeric, 6) AS lat",
              "    round(ST_X(ST_Centroid(b.geom))::numeric, 6) AS lon",
              "    (bs.effective_start AT TIME ZONE %s)::date AS calculated_on" % lit(TZ)]
    bcomments += [("lat", "Latitude of the centre"), ("lon", "Longitude of the centre"),
                  ("calculated_on", "Date of the calculation")]
    statements.append("""CREATE VIEW ai.blackspots AS SELECT
%s
FROM black_spots_blackspot b
    JOIN black_spots_blackspotset bs ON bs.uuid = b.black_spot_set_id
WHERE bs.uuid = (SELECT uuid FROM black_spots_blackspotset WHERE record_type_id = %s
                 ORDER BY effective_start DESC LIMIT 1)""" % (",\n".join(bcols), lit(rt)))
    statements.append("COMMENT ON VIEW ai.blackspots IS 'Black spots (high-risk road segments), latest calculation'")
    comments += [("blackspots", n, t) for n, t in bcomments]

    for view, col, title in comments:
        statements.append("COMMENT ON COLUMN ai.%s.%s IS %s" % (view, col, lit(title)))
    return statements, ["crashes"] + [v for _, _, v, _ in sections] + ["blackspots"]


def setup():
    conn = admin_connect()
    cur = conn.cursor()
    statements, views = build_views(cur)

    if os.path.exists(SECRETS_FILE):
        with open(SECRETS_FILE) as f:
            password = json.load(f)["password"]
    else:
        password = secrets.token_urlsafe(18)

    cur.execute("CREATE SCHEMA IF NOT EXISTS ai")
    cur.execute("SELECT table_name FROM information_schema.views WHERE table_schema = 'ai'")
    for (name,) in cur.fetchall():
        cur.execute("DROP VIEW IF EXISTS ai.%s CASCADE" % pgsql.Identifier(name).as_string(conn))
    cur.execute(FUNCTIONS_SQL)
    for st in statements:
        cur.execute(st)

    cur.execute("SELECT 1 FROM pg_roles WHERE rolname = %s", (AI_USER,))
    if cur.fetchone():
        cur.execute("ALTER ROLE driver_ai WITH LOGIN NOSUPERUSER NOCREATEDB NOCREATEROLE NOREPLICATION "
                    "PASSWORD %s", (password,))
    else:
        cur.execute("CREATE ROLE driver_ai WITH LOGIN NOSUPERUSER NOCREATEDB NOCREATEROLE NOREPLICATION "
                    "PASSWORD %s", (password,))
    cur.execute("ALTER ROLE driver_ai SET default_transaction_read_only = on")
    cur.execute("ALTER ROLE driver_ai SET statement_timeout = '15s'")
    cur.execute("REVOKE ALL ON ALL TABLES IN SCHEMA public FROM driver_ai")
    cur.execute("REVOKE CREATE ON SCHEMA public FROM PUBLIC")
    cur.execute("REVOKE ALL ON SCHEMA ai FROM PUBLIC")
    cur.execute("GRANT USAGE ON SCHEMA ai TO driver_ai")
    cur.execute("GRANT SELECT ON ALL TABLES IN SCHEMA ai TO driver_ai")
    conn.commit()
    conn.close()

    with open(SECRETS_FILE, "w") as f:
        json.dump({"user": AI_USER, "password": password}, f)
    os.chmod(SECRETS_FILE, 0o600)

    # ---- check the read-only user
    conn = ai_connect()
    cur = conn.cursor()
    print("Views created in schema 'ai':")
    for v in views:
        cur.execute("SELECT count(*) FROM ai.%s" % v)
        print("  ai.%-14s %6d rows" % (v, cur.fetchone()[0]))
    conn.rollback()
    checks = [("UPDATE grout_record SET archived = archived WHERE false", "change records"),
              ("DELETE FROM ai.crashes WHERE false", "delete through the views"),
              ("CREATE TABLE ai_check (x int)", "create tables"),
              ("SELECT count(*) FROM grout_record", "read DRIVER's tables directly")]
    ok = True
    for statement, what in checks:
        # switch off the read-only default first, to prove that permissions alone block it
        conn2 = ai_connect(readonly=False)
        conn2.autocommit = True
        try:
            conn2.cursor().execute("SET default_transaction_read_only = off")
            conn2.cursor().execute(statement)
            print("  WARNING: the AI user can %s" % what)
            ok = False
        except psycopg2.Error:
            print("  ok: the AI user cannot %s" % what)
        finally:
            conn2.close()
    conn.close()
    print("\nRead-only check %s." % ("passed" if ok else "FAILED"))


# =================================================================== database access (read-only)
def ai_connect(readonly=True):
    try:
        with open(SECRETS_FILE) as f:
            creds = json.load(f)
    except OSError:
        sys.exit("Run 'driver-ai setup' first.")
    try:
        conn = psycopg2.connect(host=DB_HOST, port=DB_PORT, dbname=DB_NAME, user=creds["user"],
                                password=creds["password"], connect_timeout=10)
    except psycopg2.OperationalError as e:
        raise PlanError(db_help(e))
    if readonly:
        conn.set_session(readonly=True)
    return conn


class Dictionary:
    """What the model is told about the data: tables, fields, values and examples."""

    def __init__(self, conn):
        cur = conn.cursor()
        cur.execute("""SELECT c.table_name, c.column_name, c.data_type,
                              col_description(('ai.' || quote_ident(c.table_name))::regclass, c.ordinal_position),
                              obj_description(('ai.' || quote_ident(c.table_name))::regclass, 'pg_class')
                       FROM information_schema.columns c
                       WHERE c.table_schema = 'ai' ORDER BY c.table_name = 'crashes' DESC,
                             c.table_name = 'blackspots', c.table_name, c.ordinal_position""")
        self.tables = {}   # table -> {"about": str, "columns": {name: info}}
        for table, column, dtype, title, about in cur.fetchall():
            t = self.tables.setdefault(table, {"about": about or "", "columns": {}})
            kind = ("number" if dtype in ("numeric", "integer", "bigint", "double precision", "real")
                    else "date" if dtype == "date" else "datetime" if dtype.startswith("timestamp")
                    else "id" if dtype == "uuid" or column in ("item_id",) or (title or "").startswith("link:")
                    else "text")
            if (title or "").startswith("Boundary: "):
                title, is_boundary = title[10:], True
            else:
                is_boundary = False
            info = {"name": column, "title": title or column, "kind": kind, "values": None, "boundary": is_boundary,
                    "hidden": column in HIDDEN_COLUMNS or column.startswith("n_")}
            if (title or "").startswith("link:"):
                info["link"] = title[5:]
            t["columns"][column] = info
        for table, t in self.tables.items():
            for col in t["columns"].values():
                ident = "ai.%s" % table
                if col["kind"] == "text":
                    cur.execute("SELECT %s, count(*) FROM %s WHERE %s IS NOT NULL GROUP BY 1 ORDER BY 2 DESC LIMIT %d"
                                % (col["name"], ident, col["name"], MAX_CATEGORIES + 1))
                    vals = [r[0] for r in cur.fetchall()]
                    if col["name"] in TIME_COLUMNS:
                        vals.sort(key=lambda v: WEEKDAYS.index(v) if v in WEEKDAYS else v)
                    if col["name"] == "year_month" and vals:
                        col["range"] = (min(vals), max(vals))
                        continue
                    if len(vals) <= MAX_CATEGORIES:
                        col["values"] = vals
                    else:
                        col["examples"] = vals[:3]
                elif col["kind"] in ("number", "date", "datetime"):
                    cur.execute("SELECT min(%s), max(%s) FROM %s" % (col["name"], col["name"], ident))
                    col["range"] = cur.fetchone()
        conn.rollback()
        self.boundaries = [c["name"] for c in self.tables.get("crashes", {}).get("columns", {}).values()
                           if c["boundary"]]

    # ---------------------------------------------------------- lookups
    def column(self, table, name):
        return self.tables.get(table, {}).get("columns", {}).get(name)

    def fields(self, visible_only=True):
        out = []
        for table, t in self.tables.items():
            for c in t["columns"].values():
                if not visible_only or (c["kind"] != "id" and not c["hidden"]):
                    out.append("%s.%s" % (table, c["name"]))
        return out

    def find(self, test, tables=None):
        for table, t in self.tables.items():
            if tables and table not in tables:
                continue
            for c in t["columns"].values():
                if c["kind"] != "id" and test(c):
                    return table, c["name"]
        return None, None

    def value_in(self, col_ref, *wanted):
        """First of the wanted values (case-insensitive substring) that exists in a column."""
        if not col_ref or not col_ref[0]:
            return None
        vals = self.column(*col_ref).get("values") or []
        for w in wanted:
            for v in vals:
                if w.lower() in str(v).lower():
                    return v
        return vals[0] if vals else None

    def section_tables(self):
        return [t for t in self.tables if t not in ("crashes", "blackspots")]

    # ---------------------------------------------------------- prompt text
    def describe(self):
        lines = []
        for table, t in self.tables.items():
            lines.append("TABLE %s: %s" % (table, t["about"]))
            for c in t["columns"].values():
                if c["kind"] == "id" or c["hidden"]:
                    continue
                desc = "  - %s.%s (%s" % (table, c["name"], c["kind"])
                if c["boundary"]:
                    desc += ", area"
                elif c["title"] and slug(c["title"]) != c["name"]:
                    desc += ', "%s"' % c["title"]
                desc += ")"
                if c.get("values") is not None:
                    vals = c["values"]
                    desc += " values: " + ", ".join(str(v) for v in vals) if vals else " (no values yet)"
                elif c.get("examples"):
                    desc += " free text, e.g. " + "; ".join(str(v)[:40] for v in c["examples"][:2])
                elif c.get("range") and c["range"][0] is not None:
                    lo, hi = c["range"]
                    desc += " from %s to %s" % (fmt(lo), fmt(hi))
                lines.append(desc)
        return "\n".join(lines)

    def examples(self, today):
        """Worked examples, built from the real field names and values."""
        sev = self.find(lambda c: c.get("values") and any(str(v).lower() == "fatal" for v in c["values"]),
                        ["crashes"])
        killed = self.find(lambda c: c["kind"] == "number" and re.search("kill|dead|death|fatal", c["name"]),
                           ["crashes"])
        bnd = ("crashes", self.boundaries[0]) if self.boundaries else (None, None)
        vt = self.find(lambda c: c.get("values") and any(re.search("bus|motor|truck|car", str(v), re.I)
                                                         for v in c["values"]), self.section_tables())
        role = self.find(lambda c: c.get("values") and any(re.search("pedestrian", str(v), re.I)
                                                           for v in c["values"]), self.section_tables())
        injury = self.find(lambda c: "injur" in c["name"] and c.get("values"), self.section_tables())
        weather = ("crashes", "weather") if self.column("crashes", "weather") else (None, None)
        last_month_end = today.replace(day=1) - dt.timedelta(days=1)
        last_month = (last_month_end.replace(day=1).isoformat(), last_month_end.isoformat())
        y = today.year - 1

        def plan(table, action="list", filters=(), date=(None, None), group=(), sort=None, desc=False,
                 limit=None, metric=None):
            return {"table": table, "action": action, "metric_field": metric,
                    "filters": [{"field": f, "op": o, "value": v} for f, o, v in filters],
                    "date_from": date[0], "date_to": date[1], "group_by": list(group),
                    "sort_by": sort, "sort_desc": desc, "limit": limit}

        ex = []
        if sev[0] and bnd[0]:
            ex.append(("show fatal crashes in %s in %d" % (self.value_in(bnd, "dhaka"), y),
                       plan("crashes", filters=[("crashes." + sev[1], "=", self.value_in(sev, "fatal")),
                                                ("crashes." + bnd[1], "=", self.value_in(bnd, "dhaka"))],
                            date=("%d-01-01" % y, "%d-12-31" % y))))
        ex.append(("how many crashes per month in %d" % today.year,
                   plan("crashes", "count", date=("%d-01-01" % today.year, today.isoformat()),
                        group=["crashes.year_month"])))
        if sev[0]:
            ex.append(("number of crashes by severity", plan("crashes", "count", group=["crashes." + sev[1]])))
        if vt[0] and killed[0]:
            ex.append(("%s crashes where someone died" % self.value_in(vt, "bus").lower(),
                       plan("crashes", filters=[("%s.%s" % vt, "=", self.value_in(vt, "bus")),
                                                ("crashes." + killed[1], ">", 0)])))
        if role[0] and injury[0]:
            ex.append(("pedestrians killed last month",
                       plan(role[0], filters=[("%s.%s" % role, "=", self.value_in(role, "pedestrian")),
                                              ("%s.%s" % injury, "=", self.value_in(injury, "fatal", "kill"))],
                            date=last_month)))
        if "blackspots" in self.tables:
            ex.append(("top 5 black spots", plan("blackspots", sort="blackspots.severity_score", desc=True,
                                                 limit=5)))
        if killed[0] and bnd[0]:
            ex.append(("total deaths per %s" % bnd[1].rstrip("s"),
                       plan("crashes", "sum", metric="crashes." + killed[1], group=["crashes." + bnd[1]])))
        if vt[0]:
            ex.append(("how many crashes involved a %s in %d" % (self.value_in(vt, "truck").lower(), y),
                       plan("crashes", "count", filters=[("%s.%s" % vt, "=", self.value_in(vt, "truck"))],
                            date=("%d-01-01" % y, "%d-12-31" % y))))
        if role[0] and injury[0] and bnd[0]:
            sex = self.find(lambda c: c.get("values") and any(str(v).lower() == "male" for v in c["values"]),
                            [role[0]])
            f = [("%s.%s" % role, "=", self.value_in(role, "driver"))]
            if sex[0]:
                f.insert(0, ("%s.%s" % sex, "=", self.value_in(sex, "male")))
            f += [("%s.%s" % injury, "=", self.value_in(injury, "serious")),
                  ("crashes." + bnd[1], "=", self.value_in(bnd, "sylhet"))]
            ex.append(("%sdrivers seriously injured in %s" % ("male " if sex[0] else "",
                                                               self.value_in(bnd, "sylhet")),
                       plan(role[0], filters=f)))
        if killed[0] and bnd[0]:
            ex.append(("average deaths per crash in each %s" % bnd[1].rstrip("s"),
                       plan("crashes", "average", metric="crashes." + killed[1], group=["crashes." + bnd[1]])))
        if vt[0]:
            ex.append(("which vehicle types are involved most",
                       plan("crashes", "count", group=["%s.%s" % vt], sort="count", desc=True)))
        if weather[0] and self.column("crashes", "weather").get("values"):
            ex.append(("crashes in rainy weather during the monsoon of %d" % y,
                       plan("crashes", filters=[("crashes.weather", "=", self.value_in(weather, "rain"))],
                            date=("%d-06-01" % y, "%d-09-30" % y))))
        ex.append(("crashes on Airport Road at night",
                   plan("crashes", filters=[("crashes.road", "contains", "Airport Road")] +
                        ([("crashes.light", "=", self.value_in(("crashes", "light"), "night"))]
                         if self.column("crashes", "light") and self.column("crashes", "light").get("values")
                         else []))))
        ex = [(q, [p]) for q, p in ex]
        # questions with several parts: one query per part
        if role[0] and injury[0]:
            f = [("%s.%s" % role, "=", self.value_in(role, "pedestrian")),
                 ("%s.%s" % injury, "=", self.value_in(injury, "fatal", "kill"))]
            d = ("%d-01-01" % y, "%d-12-31" % y)
            ex.append(("how many pedestrians were killed in %d and where did they die" % y,
                       [plan(role[0], "count", filters=f, date=d), plan(role[0], filters=f, date=d)]))
        if sev[0] and bnd[0]:
            f = [("crashes." + sev[1], "=", self.value_in(sev, "fatal")),
                 ("crashes." + bnd[1], "=", self.value_in(bnd, "dhaka"))]
            ex.append(("number of fatal crashes in %s and on which roads they happened" % self.value_in(bnd, "dhaka"),
                       [plan("crashes", "count", filters=f),
                        plan("crashes", "count", filters=f, group=["crashes.road"], sort="count", desc=True)]))
        return ex


def fmt(v):
    if isinstance(v, (dt.datetime,)):
        return v.strftime("%Y-%m-%d %H:%M")
    if isinstance(v, dt.date):
        return v.isoformat()
    if isinstance(v, decimal.Decimal):
        return int(v) if v == v.to_integral_value() else float(v)
    if isinstance(v, uuid.UUID):
        return str(v)
    return v


# =================================================================== synonyms
SYNONYMS = {
    # word in a question -> words to look for among the column's values
    "motorbike": ["motorcycle", "motorbike"], "bike": ["motorcycle"], "motor cycle": ["motorcycle"],
    "cng": ["cng", "auto"], "auto rickshaw": ["auto", "cng"], "autorickshaw": ["auto", "cng"],
    "three wheeler": ["cng", "auto", "tricycle"], "tempo": ["cng", "auto", "tricycle"],
    "lorry": ["truck"], "trucks": ["truck"], "buses": ["bus"], "cars": ["car"], "cycle": ["bicycle"],
    "killed": ["fatal", "killed", "dead"], "died": ["fatal", "killed", "dead"], "death": ["fatal", "killed"],
    "deaths": ["fatal", "killed"], "dead": ["fatal", "killed"], "deadly": ["fatal"],
    "injured": ["injury", "injured", "serious", "minor"], "hurt": ["injury", "minor"],
    "property damage": ["property"], "damage only": ["property"], "pdo": ["property"],
    "rainy": ["rain"], "raining": ["rain"], "foggy": ["fog"], "windy": ["wind"],
    "evening": ["dusk"], "dark": ["night"], "nighttime": ["night"], "morning": ["day", "dawn"],
    "chittagong": ["chattogram"], "ctg": ["chattogram"], "barisal": ["barishal"], "jessore": ["jashore"],
    "comilla": ["cumilla"], "bogra": ["bogura"],
    "man": ["male"], "men": ["male"], "woman": ["female"], "women": ["female"],
    "walker": ["pedestrian"], "pedestrians": ["pedestrian"], "drivers": ["driver"], "passengers": ["passenger"],
    "ঢাকা": ["dhaka"], "চট্টগ্রাম": ["chattogram"], "সিলেট": ["sylhet"], "খুলনা": ["khulna"],
    "রাজশাহী": ["rajshahi"], "বরিশাল": ["barishal"], "রংপুর": ["rangpur"], "ময়মনসিংহ": ["mymensingh"],
}


def match_value(value, values):
    """Map a value from the model onto one of the column's real values."""
    v = str(value).strip()
    low = v.lower()
    for x in values:
        if str(x).lower() == low:
            return x
    negative = re.compile(r"\b(not?|non|none|un\w+)\b", re.I)
    for w in SYNONYMS.get(low, []):
        for x in sorted(values, key=lambda x: bool(negative.search(str(x)))):
            if w in str(x).lower() and (not negative.search(str(x)) or negative.search(low)):
                return x
    for x in values:
        xl = str(x).lower()
        if low and (xl.startswith(low) or low in xl.split()):
            return x
    close = difflib.get_close_matches(low, [str(x).lower() for x in values], n=1, cutoff=0.75)
    if close:
        return next(x for x in values if str(x).lower() == close[0])
    return None


# =================================================================== plan -> SQL
OPS = {"=", "!=", "in", "not in", "contains", ">", ">=", "<", "<="}
OP_ALIASES = {"==": "=", "eq": "=", "equals": "=", "is": "=", "<>": "!=", "ne": "!=", "not": "!=",
              "like": "contains", "ilike": "contains", "includes": "contains", "has": "contains",
              "gt": ">", "gte": ">=", "ge": ">=", "lt": "<", "lte": "<=", "le": "<=", "not_in": "not in"}
TIME_COLUMNS = {"date", "year", "month", "year_month", "weekday", "hour", "occurred_at", "calculated_on"}
HIDDEN_COLUMNS = {"lat", "lon"}   # shown in results, but not offered to the model
WEEKDAYS = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"]


def plan_schema(d):
    return {"type": "object",
            "properties": {"queries": {"type": "array", "items": single_plan_schema(d), "minItems": 1,
                                       "maxItems": 3}},
            "required": ["queries"]}


def single_plan_schema(d):
    fields = d.fields()
    value = {"anyOf": [{"type": "string"}, {"type": "number"},
                       {"type": "array", "items": {"anyOf": [{"type": "string"}, {"type": "number"}]}}]}
    return {
        "type": "object",
        "properties": {
            "table": {"type": "string", "enum": list(d.tables)},
            "action": {"type": "string", "enum": ["list", "count", "sum", "average"]},
            "metric_field": {"anyOf": [{"type": "string", "enum": fields}, {"type": "null"}]},
            "filters": {"type": "array", "items": {
                "type": "object",
                "properties": {"field": {"type": "string", "enum": fields},
                               "op": {"type": "string", "enum": sorted(OPS)},
                               "value": value},
                "required": ["field", "op", "value"]}},
            "date_from": {"anyOf": [{"type": "string"}, {"type": "null"}]},
            "date_to": {"anyOf": [{"type": "string"}, {"type": "null"}]},
            "group_by": {"type": "array", "items": {"type": "string", "enum": fields}},
            "sort_by": {"anyOf": [{"type": "string", "enum": fields + ["count"]}, {"type": "null"}]},
            "sort_desc": {"type": "boolean"},
            "limit": {"anyOf": [{"type": "integer"}, {"type": "null"}]},
        },
        "required": ["table", "action", "filters", "date_from", "date_to", "group_by", "sort_by",
                     "sort_desc", "limit", "metric_field"],
    }


def parse_date(s):
    if not s:
        return None
    s = str(s).strip()
    for f in ("%Y-%m-%d", "%Y/%m/%d", "%d/%m/%Y", "%Y-%m", "%Y"):
        try:
            d = dt.datetime.strptime(s, f).date()
            return d
        except ValueError:
            pass
    raise PlanError("'%s' is not a date (use YYYY-MM-DD)" % s)


class Query:
    """A checked plan, ready to run. Built only from known tables and columns."""

    def __init__(self, d, plan):
        self.d, self.warnings = d, []
        if not isinstance(plan, dict):
            raise PlanError("The plan is not a JSON object")
        self.table = plan.get("table") or "crashes"
        if self.table not in d.tables:
            raise PlanError("Unknown table '%s'" % self.table)
        self.action = (plan.get("action") or "list").lower()
        if self.action not in ("list", "count", "sum", "average"):
            raise PlanError("Unknown action '%s'" % self.action)
        self.has_crash = self.table != "blackspots"

        # filters (same field and '=' are combined into 'in')
        self.filters = []
        merged = {}
        for f in plan.get("filters") or []:
            if not isinstance(f, dict):
                continue
            table, col = self.resolve(f.get("field"))
            op = str(f.get("op") or "=").lower().strip()
            op = OP_ALIASES.get(op, op)
            if op not in OPS:
                raise PlanError("Unknown operator '%s'" % op)
            value = f.get("value")
            if isinstance(value, list) and op == "=":
                op = "in"
            if op == "!=" and isinstance(value, list):
                op = "not in"
            if op in ("in", "not in") and not isinstance(value, list):
                op = "=" if op == "in" else "!="
            if op == "=" and (table, col) in merged:
                merged[(table, col)]["value"].append(value)
                merged[(table, col)]["op"] = "in"
                continue
            item = {"table": table, "col": col, "op": op, "value": [value] if op == "=" else value}
            if op == "=":
                merged[(table, col)] = item
            self.filters.append(item)
        for item in list(self.filters):
            if item["op"] in ("=", "in") and isinstance(item["value"], list) and len(item["value"]) == 1:
                item["value"], item["op"] = item["value"][0], "="
            kind = d.column(item["table"], item["col"])["kind"]
            if kind in ("date", "datetime") and item["op"] in ("=", "in") and isinstance(item["value"], list) \
                    and len(item["value"]) == 2:
                # two dates given as a list mean a range
                lo, hi = sorted(parse_date(v) for v in item["value"])
                self.filters.remove(item)
                self.filters += [dict(item, op=">=", value=lo), dict(item, op="<=", value=hi)]
        for item in self.filters:
            self.check_filter(item)

        # dates
        self.date_from, self.date_to = parse_date(plan.get("date_from")), parse_date(plan.get("date_to"))
        if plan.get("date_to") and re.fullmatch(r"\d{4}", str(plan.get("date_to")).strip()):
            self.date_to = dt.date(self.date_to.year, 12, 31)
        if self.date_from and self.date_to and self.date_from > self.date_to:
            self.date_from, self.date_to = self.date_to, self.date_from
        if (self.date_from or self.date_to) and not self.has_crash:
            self.warnings.append("Black spots have no crash dates; the date range was ignored.")
            self.date_from = self.date_to = None

        # grouping
        self.group = []
        self.join_section = None
        for g in plan.get("group_by") or []:
            table, col = self.resolve(g)
            if table == self.table or (table == "crashes" and self.has_crash):
                self.group.append((table, col))
            elif self.table == "crashes" and table in d.section_tables() and self.join_section in (None, table):
                self.join_section = table
                self.group.append((table, col))
            else:
                raise PlanError("Cannot group %s by %s.%s" % (self.table, table, col))
        numeric_groups = self.group and all(
            d.column(t, c)["kind"] == "number" and c not in TIME_COLUMNS for t, c in self.group)
        if numeric_groups and self.action in ("list", "count") and plan.get("limit"):
            # "top 5 black spots" sometimes comes back grouped by the score: list and sort instead
            plan = dict(plan, sort_by="%s.%s" % self.group[0], sort_desc=True)
            self.group, self.action, self.join_section = [], "list", None
        if self.group and self.action == "list":
            self.action = "count"
        self.metric = None
        if self.action in ("sum", "average"):
            m = plan.get("metric_field")
            if not m:
                raise PlanError("Which number should be added up? (metric_field is missing)")
            table, col = self.resolve(m)
            if d.column(table, col)["kind"] != "number":
                raise PlanError("%s.%s is not a number" % (table, col))
            if table != self.table and not (table == "crashes" and self.has_crash):
                raise PlanError("Cannot add up %s.%s for %s" % (table, col, self.table))
            if self.join_section:
                raise PlanError("Sums cannot be grouped by %s fields; group by crash fields instead"
                                % self.join_section)
            self.metric = (table, col)

        # sorting and limit
        self.sort = None
        s = plan.get("sort_by")
        if s and s != "count":
            table, col = self.resolve(s)
            if self.action == "list" and (table == self.table or (table == "crashes" and self.has_crash)):
                self.sort = (table, col)
            elif self.action != "list" and (table, col) in self.group:
                self.sort = (table, col)
        elif s == "count" and self.action != "list":
            self.sort = "value"
        self.sort_desc = bool(plan.get("sort_desc"))
        lim = plan.get("limit")
        self.limit = max(1, min(int(lim), 1000)) if isinstance(lim, (int, float)) and lim else None

    # ------------------------------------------------------------------
    def resolve(self, field):
        """'table.column' (or just 'column') -> (table, column), only for known fields."""
        if not field or not isinstance(field, str):
            raise PlanError("A field name is missing")
        field = field.strip()
        if "." in field:
            table, col = field.split(".", 1)
            c = self.d.column(table, col)
            if c and c["kind"] != "id":
                return table, col
            field = col
        order = [self.table, "crashes"] + [t for t in self.d.tables if t not in (self.table, "crashes")]
        for table in order:
            c = self.d.column(table, field)
            if c and c["kind"] != "id":
                return table, field
        raise PlanError("Unknown field '%s'" % field)

    def check_filter(self, item):
        c = self.d.column(item["table"], item["col"])
        if item["table"] == "blackspots" and self.table != "blackspots":
            raise PlanError("Black spot fields can only be used when listing black spots")
        if self.table == "blackspots" and item["table"] != "blackspots":
            raise PlanError("Black spots can only be filtered by their own fields (%s)" %
                            ", ".join(n for n, x in self.d.tables["blackspots"]["columns"].items()
                                      if x["kind"] != "id"))
        values = item["value"] if isinstance(item["value"], list) else [item["value"]]
        if c["kind"] == "number":
            if item["op"] == "contains":
                raise PlanError("'contains' cannot be used with the number %s" % item["col"])
            try:
                values = [int(float(v)) if float(v).is_integer() else float(v) for v in values]
            except (TypeError, ValueError):
                raise PlanError("%s needs a number, not %s" % (item["col"], values))
        elif c["kind"] in ("date", "datetime"):
            values = [parse_date(v) for v in values]
        elif item["op"] != "contains" and c.get("values") is not None:
            fixed = []
            for v in values:
                m = match_value(v, c["values"])
                if m is None:
                    raise PlanError("'%s' is not a value of %s. Possible values: %s" % (
                        v, c["title"], ", ".join(str(x) for x in c["values"][:25])))
                if str(m) != str(v):
                    self.warnings.append("'%s' was read as '%s'" % (v, m))
                fixed.append(m)
            values = fixed
        else:
            values = [str(v) for v in values]
            if item["op"] in (">", ">=", "<", "<="):
                raise PlanError("%s is text; '%s' cannot be used" % (item["col"], item["op"]))
        item["value"] = values if item["op"] in ("in", "not in") else values[0]

    # ------------------------------------------------------------------
    def describe(self):
        """Human-readable summary of what will be shown."""
        if self.metric:
            what = "%s %s (%s)" % ("Total" if self.action == "sum" else "Average",
                                   self.d.column(*self.metric)["title"].lower(), self.table)
        else:
            what = "%s %s" % ("List of" if self.action == "list" else "Number of", self.table)
        parts = [what]
        if self.group:
            parts.append("by " + ", ".join(self.d.column(t, c)["title"] for t, c in self.group))
        for f in self.filters:
            title = self.d.column(f["table"], f["col"])["title"]
            if f["table"] != self.table:
                title = "%s %s" % (f["table"], title)
            v = f["value"]
            v = " or ".join(str(fmt(x)) for x in v) if isinstance(v, list) else str(fmt(v))
            op = {"=": "=", "!=": "≠", "in": "=", "not in": "≠", "contains": "contains"}.get(f["op"], f["op"])
            parts.append("%s %s %s" % (title, op, v))
        if self.date_from or self.date_to:
            parts.append("%s to %s" % (self.date_from or "start", self.date_to or "today"))
        if self.limit:
            parts.append("top %d" % self.limit)
        return " · ".join(parts)

    def condition(self, alias, f, params):
        col = pgsql.SQL("{}.{}").format(pgsql.Identifier(alias), pgsql.Identifier(f["col"]))
        kind = self.d.column(f["table"], f["col"])["kind"]
        op, v = f["op"], f["value"]
        if op == "contains":
            params.append("%" + str(v).replace("%", r"\%").replace("_", r"\_") + "%")
            return pgsql.SQL("{} ILIKE %s").format(col)
        if op in ("in", "not in"):
            params.append(tuple(v))
            if kind == "text":
                params[-1] = tuple(str(x).lower() for x in v)
                return pgsql.SQL("lower({}) " + op.upper() + " %s").format(col)
            return pgsql.SQL("{} " + op.upper() + " %s").format(col)
        params.append(str(v).lower() if kind == "text" and op in ("=", "!=") else v)
        if kind == "text" and op in ("=", "!="):
            return pgsql.SQL("lower({}) " + op + " %s").format(col)
        return pgsql.SQL("{} " + op + " %s").format(col)

    def where(self, params):
        conds = []
        if self.date_from:
            params.append(self.date_from)
            conds.append(pgsql.SQL("c.date >= %s"))
        if self.date_to:
            params.append(self.date_to)
            conds.append(pgsql.SQL("c.date <= %s"))
        main = "c" if self.table == "crashes" else "t"
        for f in self.filters:
            if f["table"] == self.table:
                conds.append(self.condition(main, f, params))
            elif f["table"] == "crashes":
                conds.append(self.condition("c", f, params))
            else:   # a field of another section: EXISTS
                link = None
                if self.table != "crashes":
                    link = next((n for n, x in self.d.tables[self.table]["columns"].items()
                                 if x.get("link") == f["table"]), None)
                sub_params = []
                cond = self.condition("x", f, sub_params)
                if link:
                    q = pgsql.SQL("EXISTS (SELECT 1 FROM ai.{} x WHERE x.record_id = c.record_id "
                                  "AND x.item_id = t.{} AND {})").format(
                        pgsql.Identifier(f["table"]), pgsql.Identifier(link), cond)
                else:
                    q = pgsql.SQL("EXISTS (SELECT 1 FROM ai.{} x WHERE x.record_id = c.record_id AND {})").format(
                        pgsql.Identifier(f["table"]), cond)
                params.extend(sub_params)
                conds.append(q)
        if not conds:
            return pgsql.SQL("")
        return pgsql.SQL(" WHERE ") + pgsql.SQL(" AND ").join(conds)

    def from_clause(self):
        if self.table == "crashes":
            q = pgsql.SQL(" FROM ai.crashes c")
            if self.join_section:
                q += pgsql.SQL(" JOIN ai.{} j ON j.record_id = c.record_id").format(
                    pgsql.Identifier(self.join_section))
            return q
        if self.table == "blackspots":
            return pgsql.SQL(" FROM ai.blackspots t")
        return pgsql.SQL(" FROM ai.{} t JOIN ai.crashes c ON c.record_id = t.record_id").format(
            pgsql.Identifier(self.table))

    def alias(self, table):
        if table == self.join_section:
            return "j"
        if self.table == "crashes":
            return "c"
        return "t" if table == self.table else "c"

    def key(self):
        if self.table == "crashes":
            return pgsql.SQL("c.record_id")
        if self.table == "blackspots":
            return pgsql.SQL("t.blackspot_id")
        return pgsql.SQL("(t.record_id::text || coalesce(t.item_id, ''))")

    def list_columns(self):
        """(alias, column, header) for a list result."""
        cols = []
        crash = self.d.tables["crashes"]["columns"]
        skip = {"date", "year", "month", "year_month", "weekday", "hour", "record_id"}
        if self.table == "crashes":
            for name, c in crash.items():
                if name not in skip and c["kind"] != "id":
                    cols.append(("c", name, c["title"]))
            cols.append(("c", "record_id", "Record ID"))
        elif self.table == "blackspots":
            for name, c in self.d.tables["blackspots"]["columns"].items():
                if c["kind"] != "id":
                    cols.append(("t", name, c["title"]))
        else:
            cols.append(("c", "occurred_at", "Date and time"))
            for b in self.d.boundaries:
                cols.append(("c", b, crash[b]["title"]))
            sev = next((n for n, c in crash.items() if c.get("values") and
                        any(str(v).lower() == "fatal" for v in c["values"])), None)
            if sev:
                cols.append(("c", sev, "Crash " + crash[sev]["title"].lower()))
            for name in ("road", "city", "lat", "lon"):
                if name in crash:
                    cols.append(("c", name, crash[name]["title"]))
            for name, c in self.d.tables[self.table]["columns"].items():
                if c["kind"] != "id":
                    cols.append(("t", name, c["title"]))
            cols.append(("c", "record_id", "Record ID"))
        return cols

    def build(self, limit):
        """SQL for the result rows -> (sql, params, headers)."""
        params = []
        where = self.where(params)
        if self.action == "list":
            cols = self.list_columns()
            select = pgsql.SQL(", ").join(
                pgsql.SQL("{}.{}").format(pgsql.Identifier(a), pgsql.Identifier(n)) for a, n, _ in cols)
            if self.sort:
                order = pgsql.SQL("{}.{} {} NULLS LAST").format(
                    pgsql.Identifier(self.alias(self.sort[0])), pgsql.Identifier(self.sort[1]),
                    pgsql.SQL("DESC" if self.sort_desc else "ASC"))
            elif self.table == "blackspots":
                order = pgsql.SQL("t.severity_score DESC")
            else:
                order = pgsql.SQL("c.occurred_at DESC")
            q = pgsql.SQL("SELECT ") + select + self.from_clause() + where + pgsql.SQL(" ORDER BY ") + order
            headers = [h for _, _, h in cols]
        else:
            gcols = [pgsql.SQL("{}.{}").format(pgsql.Identifier(self.alias(t)), pgsql.Identifier(c))
                     for t, c in self.group]
            if self.action == "count":
                value = pgsql.SQL("count(DISTINCT ") + self.key() + pgsql.SQL(")")
                vname = "Number of " + self.table
            else:
                m = pgsql.SQL("{}.{}").format(pgsql.Identifier(self.alias(self.metric[0])),
                                              pgsql.Identifier(self.metric[1]))
                fn = "sum" if self.action == "sum" else "avg"
                value = pgsql.SQL("round(%s({})::numeric, 2)" % fn).format(m)
                vname = ("Total " if self.action == "sum" else "Average ") + \
                    self.d.column(*self.metric)["title"].lower()
            select = pgsql.SQL(", ").join(gcols + [value + pgsql.SQL(" AS value")])
            q = pgsql.SQL("SELECT ") + select + self.from_clause() + where
            if gcols:
                q += pgsql.SQL(" GROUP BY ") + pgsql.SQL(", ").join(gcols)
                time_group = all(c in TIME_COLUMNS for _, c in self.group)
                if self.sort == "value" or (self.sort is None and not time_group):
                    q += pgsql.SQL(" ORDER BY value %s NULLS LAST" % ("ASC" if self.sort == "value"
                                                                      and not self.sort_desc else "DESC"))
                else:
                    srt = self.sort if isinstance(self.sort, tuple) else self.group[0]
                    q += pgsql.SQL(" ORDER BY {}.{} {}").format(
                        pgsql.Identifier(self.alias(srt[0])), pgsql.Identifier(srt[1]),
                        pgsql.SQL("DESC" if self.sort_desc and self.sort else "ASC"))
            headers = [self.d.column(t, c)["title"] for t, c in self.group] + [vname]
        lim = self.limit or limit
        if lim:
            q += pgsql.SQL(" LIMIT %d" % lim)
        return q, params, headers

    def count_sql(self):
        params = []
        where = self.where(params)
        q = pgsql.SQL("SELECT count(DISTINCT ") + self.key() + pgsql.SQL(")") + self.from_clause() + where
        return q, params


def run_query(conn, query, limit):
    cur = conn.cursor()
    try:
        q, params, headers = query.build(limit)
        cur.execute(q, params)
        rows = [[fmt(v) for v in r] for r in cur.fetchall()]
        cq, cparams = query.count_sql()
        cur.execute(cq, cparams)
        matched = cur.fetchone()[0]
        lim = query.limit or limit
        return {"headers": headers, "rows": rows, "matched": matched, "sql": cur.mogrify(q, params).decode(),
                "complete": not (lim and len(rows) >= lim)}
    finally:
        conn.rollback()


# =================================================================== the model
SYSTEM_PROMPT = """You convert questions about a road-crash database into JSON query plans.
Never answer the question yourself and never invent fields or values: only output the plans.
Output {{"queries": [plan, ...]}}: one plan per thing the user asks (at most 3).

Today is {today} ({weekday}). Time zone: {tz}.

{tables}

HOW TO WRITE A PLAN
- table: what the user wants to see or count: crashes (default){sections}, or blackspots.
  Questions about people (pedestrians, drivers, passengers, victims, men, women, ages) use the people table.
  "How many crashes ..." always uses table crashes with action count, even when it mentions vehicles.
- action: "list" to show records; "count" for how many (add group_by for "per", "by", "each");
  "sum"/"average" with metric_field set to a number field (e.g. total deaths).
  A question that only names records ("fatal crashes in 2025", "show ...", "list ...") is "list".
  Use "count" only for "how many", "number of", "per", "by", "each", "most".
- filters: [{{"field": "table.column", "op": "=", "value": ...}}]. Use exact values from the lists above.
  op: = != in "not in" contains > >= < <=. Use "in" with a list for several values.
  Use "contains" for free-text fields (road, address, plate numbers, names).
  Filters on another table mean "crashes that have such a vehicle/person".
- date_from / date_to: YYYY-MM-DD, both inclusive, or null when no time is mentioned.
  "in 2025" = 2025-01-01 to 2025-12-31. "last year" = {last_year}-01-01 to {last_year}-12-31.
  "last month" = the previous calendar month. "last 30 days" = 30 days before today to today.
  "this year" = {year}-01-01 to today. "monsoon" or "rainy season" = June 1 to September 30.
- Months, weekdays and hours: use crashes.month (1-12), crashes.weekday (Monday...Sunday), crashes.hour (0-23).
  "at night" = crashes.light night if available, otherwise crashes.hour >= 19 or <= 5.
- "top N", "most", "highest": set sort_by (a field, or "count" for grouped results), sort_desc true and limit N.
- "where", "locations", "which places": a list (each record shows its area, road and place).
- Use null for anything not mentioned. group_by, filters: [] when not needed.

SEVERAL QUESTIONS AT ONCE
If the user asks two or three things ("how many ... and where ...", "... and also ..."), write one plan
for each part, with the same filters and dates.

FOLLOW-UP QUESTIONS
Earlier questions and your plans are part of the conversation. A follow-up ("only in Dhaka",
"what about 2024?", "now by vehicle type", "how many?", "show them", "and the injured ones?")
changes the most recent plans (your last answer): keep their table, action, filters and dates and
change only what is asked; keep every other filter. Example: after a plan for Dhaka in 2025,
"what about 2024?" = the same plan, still Dhaka, with 2024 dates. "how many?" = the same filters with
action count. A new, unrelated question starts fresh.

COMMON WORDS
bike/motorbike = Motorcycle; CNG/auto-rickshaw/tempo = CNG or Auto-rickshaw; lorry = Truck;
killed/died/deaths = Fatal (severity or injury) or the number-killed field; hurt/injured = Injury;
{extra_words}
Chittagong/CTG = Chattogram; Barisal = Barishal.

EXAMPLES (fields not shown are null, [] or false)
{examples}
"""


def build_prompt(d, today=None):
    today = today or dt.date.today()
    secs = d.section_tables()
    def compact(p):   # fields left out of an example are null / [] / false
        return {k: v for k, v in p.items() if v not in (None, [], False)}
    examples = "\n".join("Q: %s\nA: %s" % (q, json.dumps({"queries": [compact(p) for p in ps]}, ensure_ascii=False))
                         for q, ps in d.examples(today))
    extra = []
    injury = d.find(lambda c: "injur" in c["name"] and c.get("values"), d.section_tables())
    if injury[0]:
        hurt = [v for v in d.column(*injury)["values"]
                if not re.search(r"\bnot?\b|none|fatal|kill|dead", str(v), re.I)]
        if hurt:
            extra.append("injured people = %s.%s in %s" % (injury[0], injury[1], json.dumps(hurt)))
    return SYSTEM_PROMPT.format(today=today.isoformat(), weekday=today.strftime("%A"), tz=TZ,
                                tables=d.describe(), year=today.year, last_year=today.year - 1,
                                extra_words="\n".join(extra),
                                sections=(", " + ", ".join(secs)) if secs else "", examples=examples)


def ollama_chat(messages, schema):
    body = {"model": MODEL, "messages": messages, "format": schema, "stream": False,
            "keep_alive": "30m", "options": {"temperature": 0, "num_ctx": 8192}}
    if re.match(r"(qwen3|deepseek-r1|gpt-oss|magistral)", MODEL):
        body["think"] = False
    req = urllib.request.Request(OLLAMA_URL + "/api/chat", data=json.dumps(body).encode(),
                                 headers={"Content-Type": "application/json"})
    try:
        with urllib.request.urlopen(req, timeout=900) as r:
            res = json.loads(r.read().decode())
    except urllib.error.HTTPError as e:
        msg = e.read().decode()[:300]
        if "not found" in msg:
            raise PlanError("The model '%s' is not installed. Run: ollama pull %s" % (MODEL, MODEL))
        raise PlanError("Ollama error: %s" % msg)
    except urllib.error.URLError as e:
        raise PlanError("Cannot reach Ollama at %s (is it running? 'systemctl status ollama'): %s"
                        % (OLLAMA_URL, e.reason))
    except OSError as e:    # includes timeouts
        raise PlanError("The model took too long to answer (%s). Try again; if it keeps happening, "
                        "use a smaller model such as qwen2.5:1.5b." % e)
    return res["message"]["content"]


GROUP_WORDS = (r"\b(per|by|each|every|which|most|least|top|highest|lowest|worst|best|compare|comparison|"
               r"breakdown|split|distribution|where|locat\w*|places|areas|roads|divisions|districts|"
               r"monthly|yearly|weekly|daily|hourly|trend)\b")


def subject_fix(d, question, plan):
    """Small models tend to list crashes even when the question is about people or vehicles.
    If the question never mentions crashes and every filter/grouping outside the crash table
    belongs to one section (e.g. people), use that section instead."""
    if not isinstance(plan, dict):
        return plan
    # "where did they die?", "tell me the locations": list the records (each shows its place)
    if re.search(r"\b(where|locat\w*|places?)\b", question, re.I) and \
            not re.search(r"\b(how many|number|count|total|sum|average|per|by|each|most|which)\b", question, re.I):
        plan = dict(plan, action="list", group_by=[], metric_field=None, limit=None, sort_by=None)
    # grouping only when the question asks for it ("per", "by", "which", "most", "where", ...)
    if plan.get("group_by") and not re.search(GROUP_WORDS, question, re.I):
        plan = dict(plan, group_by=[], limit=None, sort_by=None, sort_desc=False)
    if plan.get("table") != "crashes":
        return plan
    if plan.get("action") in ("list", "count"):
        plan = dict(plan, metric_field=None)     # only used by sum/average
    if re.search(r"crash|accident|incident|collision|record|case|event", question, re.I):
        return plan
    used = set()
    for f in plan.get("filters") or []:
        if isinstance(f, dict) and isinstance(f.get("field"), str) and "." in f["field"]:
            used.add(f["field"].split(".", 1)[0])
    for g in plan.get("group_by") or []:
        if isinstance(g, str) and "." in g:
            used.add(g.split(".", 1)[0])
    sections = [t for t in used if t in d.section_tables()]
    if len(sections) == 1 and plan.get("action") in ("list", "count"):
        plan = dict(plan, table=sections[0])
        if plan.get("sort_by") and str(plan["sort_by"]).split(".", 1)[0] not in (sections[0], "crashes", "count"):
            plan["sort_by"] = None
    return plan


def summarize(d, query, res):
    """One plain sentence about a result, written by code from the real numbers (never by the model)."""
    rows, table = res["rows"], query.table
    if not rows:
        return "No matching %s." % table
    if query.action == "list":
        text = "Found %s %s." % (res["matched"], table)
        if len(rows) < res["matched"]:
            text += " The first %d are shown; the export has all of them." % len(rows)
        if table == "blackspots" and rows:
            h = res["headers"]
            first = dict(zip(h, rows[0]))
            places = [k for k in h if k == "Main road" or (d.column("blackspots", slug(k)) or {}).get("boundary")]
            where = ", ".join(str(first[k]) for k in places if first.get(k))
            text += " The most dangerous one is %s(risk score %s)." % ((where + " ") if where else "",
                                                                         first.get(h[0]))
        return text
    value_name = res["headers"][-1]
    if not query.group:
        v = rows[0][-1]
        if query.action == "count":
            return "%s %s." % (v, table if v != 1 else table.rstrip("s"))
        return "%s: %s." % (value_name, v)
    valid = [r for r in rows if r[-1] is not None]
    if not valid:
        return "No matching %s." % table
    label = lambda r: " / ".join(str(x) for x in r[:-1])
    if all(c in TIME_COLUMNS for _, c in query.group):
        hi = max(valid, key=lambda r: r[-1])
        lo = min(valid, key=lambda r: r[-1])
        text = "Highest: %s (%s). Lowest: %s (%s)." % (label(hi), hi[-1], label(lo), lo[-1])
    else:
        top = sorted(valid, key=lambda r: r[-1], reverse=True)[:3]
        text = "Most: " + ", ".join("%s (%s)" % (label(r), r[-1]) for r in top) + "."
    if query.action == "count" and not query.join_section:
        text += " Total: %s %s." % (res["matched"], table)
    elif query.action == "sum" and res.get("complete") and not query.limit:
        total = sum(r[-1] for r in valid)
        text += " Total: %s." % (int(total) if float(total).is_integer() else round(total, 2))
    return text


MAX_HISTORY = 3

SPLIT_RE = re.compile(
    r"\s*[,;]?\s*\band (?:also|then|tell me|show me|give me|list)\b[,:]?\s*"     # "... and also, tell me ..."
    r"|\s*[,;.]\s*(?:also|then|plus)\b[,:]?\s*"                                  # "..., also ..."
    r"|\s+and\s+(?=(?:where|which|what|when|who|how many|how much)\b)"            # "... and where ..."
    r"|\s*[?;]\s+(?=\S)", re.I)                                                   # "...? ..."


def split_question(question):
    """A question with several parts -> the parts (asked one after the other, as follow-ups).
    Small models answer one thing at a time much more reliably than several at once."""
    parts = [p.strip(" ,.;") for p in SPLIT_RE.split(question)]
    parts = [p for p in parts if p]
    if len(parts) < 2 or any(len(p.split()) < 2 for p in parts):
        return [question]
    return parts[:3]


def history_messages(history):
    """Earlier turns of the conversation (questions and the plans that answered them)."""
    msgs = []
    for turn in (history or [])[-MAX_HISTORY:]:
        if not isinstance(turn, dict):
            continue
        q = str(turn.get("question") or "")[:500]
        plans = [p for p in (turn.get("plans") or []) if isinstance(p, dict)][:3]
        if not q or not plans:
            continue
        clean = [{k: v for k, v in p.items() if not str(k).startswith("_")} for p in plans]
        text = json.dumps({"queries": clean}, ensure_ascii=False)
        if len(text) > 4000:
            continue
        msgs += [{"role": "user", "content": q}, {"role": "assistant", "content": text}]
    return msgs


class Assistant:
    def __init__(self):
        self.refresh()

    def refresh(self):
        conn = ai_connect()
        try:
            self.d = Dictionary(conn)
        finally:
            conn.close()
        self.schema = plan_schema(self.d)
        self.prompt = build_prompt(self.d)
        self.prompt_day = dt.date.today()

    def plan(self, question, history=None):
        """Question (+ earlier turns) -> list of (plan, Query), plus notes about parts that failed."""
        if self.prompt_day != dt.date.today():
            self.prompt = build_prompt(self.d)
            self.prompt_day = dt.date.today()
        messages = ([{"role": "system", "content": self.prompt}] + history_messages(history) +
                    [{"role": "user", "content": question}])
        errors = []
        for attempt in range(2):
            content = ollama_chat(messages, self.schema)
            try:
                data = json.loads(content)
            except ValueError as e:
                errors = [str(e)]
                continue
            plans = data.get("queries") if isinstance(data, dict) and "queries" in data else [data]
            done, errors, seen = [], [], {}
            for p in (plans or [])[:3]:
                try:
                    p = subject_fix(self.d, question, p)
                    query = Query(self.d, p)
                except (PlanError, TypeError, ValueError, AttributeError) as e:
                    errors.append(str(e))
                    continue
                # the same query twice (perhaps once as "top 1"): keep one, preferring the one without a limit
                key = " · ".join(sorted(x for x in query.describe().split(" · ") if not x.startswith("top ")))
                if key in seen:
                    i = seen[key]
                    if done[i][1].limit and not query.limit:
                        done[i] = (p, query)
                    continue
                seen[key] = len(done)
                done.append((p, query))
            if done:
                notes = ["One part of the question could not be answered: %s" % e for e in errors]
                return done, notes
            messages += [{"role": "assistant", "content": content},
                         {"role": "user", "content": "That plan is not valid: %s. Write corrected plans."
                          % "; ".join(errors)}]
        raise PlanError("; ".join(errors) or "The question could not be understood")

    def ask(self, question, history=None, limit=PREVIEW_ROWS):
        started = time.time()
        entry = {"time": dt.datetime.now().isoformat(timespec="seconds"), "question": question, "model": MODEL,
                 "history": len(history or [])}
        try:
            parts = split_question(question)
            pairs, notes, part_of = [], [], []
            turns = list(history or [])
            for part in parts:
                ps, ns = self.plan(part, turns)
                pairs += ps
                notes += ns
                part_of += [part if len(parts) > 1 else None] * len(ps)
                turns.append({"question": part, "plans": [p for p, _ in ps]})
            results = []
            conn = ai_connect()
            try:
                for (plan, query), part in zip(pairs, part_of):
                    res = run_query(conn, query, limit)
                    res.update({"plan": plan, "understood": query.describe(), "warnings": query.warnings,
                                "part": part,
                                "table": query.table, "action": query.action,
                                "summary": summarize(self.d, query, res)})
                    results.append(res)
            finally:
                conn.close()
            entry.update({"parts": parts, "plans": [p for p, _ in pairs], "sql": [r["sql"] for r in results],
                          "matched": [r["matched"] for r in results]})
            return {"ok": True, "question": question, "results": results, "notes": notes,
                    "plans": [p for p, _ in pairs], "seconds": round(time.time() - started, 1),
                    "driver_url": DRIVER_URL}
        except PlanError as e:
            entry["error"] = str(e)
            return {"ok": False, "question": question, "error": str(e),
                    "seconds": round(time.time() - started, 1)}
        finally:
            log(entry)

    def export(self, plan, fmt_name):
        query = Query(self.d, plan)     # the plan is checked again; it is never trusted
        conn = ai_connect()
        try:
            res = run_query(conn, query, EXPORT_ROWS)
        finally:
            conn.close()
        stamp = dt.datetime.now().strftime("%Y%m%d-%H%M")
        if fmt_name == "xlsx":
            try:
                import openpyxl
            except ImportError:
                raise PlanError("Excel export needs openpyxl: uv pip install --python "
                                "~/driver-ai-venv/bin/python openpyxl")
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Result"
            ws.append(res["headers"])
            for r in res["rows"]:
                ws.append(r)
            info = wb.create_sheet("Query")
            for k, v in (("Question", plan.get("_question", "")), ("Understood as", query.describe()),
                         ("Rows", len(res["rows"])), ("Exported", dt.datetime.now().isoformat(timespec="seconds"))):
                info.append([k, v])
            buf = io.BytesIO()
            wb.save(buf)
            return buf.getvalue(), "driver-ai-%s.xlsx" % stamp, \
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        buf = io.StringIO()
        w = csv.writer(buf)
        w.writerow(res["headers"])
        w.writerows(res["rows"])
        return ("﻿" + buf.getvalue()).encode("utf-8"), "driver-ai-%s.csv" % stamp, "text/csv; charset=utf-8"


def log(entry):
    try:
        os.makedirs(LOG_DIR, exist_ok=True)
        with open(os.path.join(LOG_DIR, "questions.jsonl"), "a") as f:
            f.write(json.dumps(entry, ensure_ascii=False, default=str) + "\n")
    except OSError:
        pass


# =================================================================== terminal output
def print_table(headers, rows, max_rows=30):
    shown = rows[:max_rows]
    widths = [min(28, max([len(str(h))] + [len(str(r[i])) for r in shown])) for i, h in enumerate(headers)]
    print("  ".join(str(h)[:w].ljust(w) for h, w in zip(headers, widths)))
    print("  ".join("-" * w for w in widths))
    for r in shown:
        print("  ".join(str("" if v is None else v)[:w].ljust(w) for v, w in zip(r, widths)))
    if len(rows) > len(shown):
        print("... %d more rows (use the web page to see and export everything)" % (len(rows) - len(shown)))


def print_answer(ans, show_sql=False, max_rows=30):
    if not ans["ok"]:
        print("Sorry: %s" % ans["error"])
        return
    for n in ans["notes"]:
        print("Note: %s" % n)
    for i, res in enumerate(ans["results"], 1):
        if len(ans["results"]) > 1:
            print("\n--- Part %d%s ---" % (i, (": " + res["part"]) if res.get("part") else ""))
        print("Understood as: %s" % res["understood"])
        for w in res["warnings"]:
            print("Note: %s" % w)
        print("Answer: %s" % res["summary"])
        if show_sql:
            print("SQL: %s" % res["sql"])
        print()
        print_table(res["headers"], res["rows"], max_rows)
    print("\n(%.1f s)" % ans["seconds"])


TEST_QUESTIONS = [
    "how many crashes are there",
    "show fatal crashes in Dhaka in 2025",
    "number of crashes per month in 2026",
    "crashes by severity",
    "how many crashes involved a motorcycle",
    "bus crashes where someone was killed",
    "pedestrians killed last year",
    "top 5 black spots",
    "total deaths per division",
    "which vehicle types are involved in the most crashes",
    "crashes on Airport Road",
    "crashes during the monsoon of 2025 in rain",
    "how many crashes happened at night",
    "crashes by hour of the day",
    "female passengers injured in Chattogram",
    "average number of injured per crash by severity",
    "crashes on Fridays",
    "list the 10 most recent fatal crashes",
    "trucks in head on collisions",
    "how many people were killed in 2026",
    "how many pedestrians were killed last year and where did they die",
]
TEST_CONVERSATION = ["fatal crashes in 2025", "only in Sylhet", "how many were there?", "what about 2024?"]


def test(questions, conversation=None):
    a = Assistant()
    ok = total = 0
    for q in questions:
        print("=" * 100)
        print("Q: %s" % q)
        ans = a.ask(q, limit=5)
        print_answer(ans, max_rows=5)
        ok += ans["ok"]
        total += 1
    if conversation:
        print("=" * 100)
        print("Follow-up questions (one conversation):")
        history = []
        for q in conversation:
            print("-" * 100)
            print("Q: %s" % q)
            ans = a.ask(q, history=history, limit=5)
            print_answer(ans, max_rows=5)
            ok += ans["ok"]
            total += 1
            if ans["ok"]:
                history.append({"question": q, "plans": ans["plans"]})
    print("=" * 100)
    print("%d of %d questions answered (check that each 'Understood as' line is right)." % (ok, total))


def chat():
    a = Assistant()
    history = []
    print("DRIVER AI assistant (model %s). Ask a question; follow-up questions use the earlier answers.\n"
          "Type 'new' to start over, 'quit' to stop." % MODEL)
    while True:
        try:
            q = input("\nYou: ").strip()
        except (EOFError, KeyboardInterrupt):
            print()
            break
        if not q:
            continue
        if q.lower() in ("quit", "exit", "q"):
            break
        if q.lower() in ("new", "reset", "clear"):
            history = []
            print("Started a new conversation.")
            continue
        ans = a.ask(q, history=history)
        print()
        print_answer(ans)
        if ans["ok"]:
            history = (history + [{"question": q, "plans": ans["plans"]}])[-MAX_HISTORY:]


def main():
    args = sys.argv[1:]
    if not args or args[0] in ("-h", "--help", "help"):
        print(__doc__)
        return
    cmd = args[0]
    if cmd == "setup":
        setup()
    elif cmd == "dictionary":
        print(build_prompt(Assistant().d))
    elif cmd == "ask":
        if len(args) < 2:
            sys.exit('Usage: driver-ai ask "your question" [--sql]')
        q = " ".join(a for a in args[1:] if a != "--sql")
        print_answer(Assistant().ask(q), show_sql="--sql" in args)
    elif cmd == "chat":
        chat()
    elif cmd == "test":
        if args[1:]:
            test(args[1:])
        else:
            test(TEST_QUESTIONS, TEST_CONVERSATION)
    elif cmd == "serve":
        import server
        server.serve(Assistant(), PORT)
    else:
        sys.exit("Unknown command '%s'. Run 'driver-ai help'." % cmd)


if __name__ == "__main__":
    try:
        main()
    except PlanError as e:
        sys.exit(str(e))
